
from io import BytesIO

from docx import Document
from pypdf import PdfReader


class ResumeParsingError(ValueError):
    """Raised when a resume cannot be parsed safely."""


def extract_pdf_text(file_content: bytes) -> str:
    try:
        reader = PdfReader(
            BytesIO(file_content),
            strict=True,
        )

        if reader.is_encrypted:
            raise ResumeParsingError(
                "Password-protected PDFs are not supported"
            )

        pages = []

        for page in reader.pages:
            page_text = page.extract_text() or ""

            if page_text.strip():
                pages.append(page_text.strip())

        return "\n\n".join(pages)

    except ResumeParsingError:
        raise

    except Exception as exc:
        raise ResumeParsingError(
            "Unable to read the PDF resume"
        ) from exc


def extract_docx_text(file_content: bytes) -> str:
    try:
        document = Document(BytesIO(file_content))
        content = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                content.append(text)

        for table in document.tables:
            for row in table.rows:
                cell_values = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                line = " | ".join(
                    value for value in cell_values if value
                )

                if line:
                    content.append(line)

        return "\n".join(content)

    except Exception as exc:
        raise ResumeParsingError(
            "Unable to read the DOCX resume"
        ) from exc


def extract_resume_text(
    file_content: bytes,
    file_type: str,
) -> str:
    if file_type == "pdf":
        text = extract_pdf_text(file_content)

    elif file_type == "docx":
        text = extract_docx_text(file_content)

    else:
        raise ResumeParsingError(
            "Unsupported resume file type"
        )

    cleaned_text = text.strip()

    if not cleaned_text:
        raise ResumeParsingError(
            "No extractable text found in resume"
        )

    return cleaned_text