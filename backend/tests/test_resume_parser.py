
from io import BytesIO

import pytest
from app.services.resume_parser import (
    ResumeParsingError,
    extract_resume_text,
)
from docx import Document
from pypdf import PdfWriter


def make_docx() -> bytes:
    document = Document()

    document.add_heading("Jane Example", level=1)
    document.add_paragraph("Python Developer")
    document.add_paragraph("FastAPI and PostgreSQL")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Education"
    table.cell(0, 1).text = "BS Computer Science"

    output = BytesIO()
    document.save(output)

    return output.getvalue()


def test_extract_docx():
    content = make_docx()

    text = extract_resume_text(
        content,
        "docx",
    )

    assert "Jane Example" in text
    assert "Python Developer" in text
    assert "FastAPI" in text
    assert "BS Computer Science" in text


def test_empty_docx():
    document = Document()
    output = BytesIO()
    document.save(output)

    with pytest.raises(
        ResumeParsingError,
        match="No extractable text",
    ):
        extract_resume_text(
            output.getvalue(),
            "docx",
        )


def test_blank_pdf():
    writer = PdfWriter()
    writer.add_blank_page(
        width=595,
        height=842,
    )

    output = BytesIO()
    writer.write(output)

    with pytest.raises(
        ResumeParsingError,
        match="No extractable text",
    ):
        extract_resume_text(
            output.getvalue(),
            "pdf",
        )


def test_unsupported_type():
    with pytest.raises(
        ResumeParsingError,
        match="Unsupported",
    ):
        extract_resume_text(
            b"example",
            "txt",
        )


def test_corrupted_pdf():
    with pytest.raises(ResumeParsingError):
        extract_resume_text(
            b"%PDF-invalid-content",
            "pdf",
        )
